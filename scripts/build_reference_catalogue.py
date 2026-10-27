"""Rebuild the plant-parameter reference catalogue from the source workbook.

Usage::

    python scripts/build_reference_catalogue.py "Power Plant Parameters English.xlsx" \
        --out data/reference/plant_parameters

The exports are *reference data*, not study input.  This script only
transcribes cells: it never converts units, fills gaps or reconciles conflicts.
Every record keeps its source sheet, its 1-based source row and the raw cell
values; interpretation and quality checks live in ``powersim.reference``.

The workbook has a few layout features that a naive one-header-row export
loses, which is why this script exists:

* ``Hydropower plants`` contains two sections (run-of-river and storage) with
  different column layouts and separate two-row headers.
* ``HPP Summary`` has a two-row header (group + sub-column).
* ``Stations by Location`` mixes plant rows with section headings and
  aggregate "Total" rows.
"""
from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path
from typing import Any

SCHEMA = "powersim.reference.table.v2"

STATION_COLUMNS = [
    ("plant_name", "Power Plant Name", None),
    ("generating_units", "Generating Units", "units x MW"),
    ("installed_capacity_mw", "Installed Capacity (MW)", "MW"),
    ("available_capacity_mw", "Available Capacity (MW)", "MW"),
    ("design_generation_gwh", "Design Generation (million kWh)", "million kWh"),
    ("year_commissioned", "Year Commissioned", "year"),
    ("regulation_type", "Regulation Type", None),
    ("river", "River", None),
    ("river_basin", "River Basin", None),
    ("region", "Region", None),
    ("city_village", "City / Village", None),
]

RESERVOIR_COLUMNS = [
    ("reservoir_name", "Reservoir name", None),
    ("included_in_model", "Included in the model?", None),
    ("max_volume_mm3", "Reservoir max volume (Mm3)", "Mm3"),
    ("active_volume_mm3", "Reservoir active volume (Mm3)", "Mm3"),
    ("min_operating_volume_mm3", "Reservoir min operating volume (Mm3)", "Mm3"),
    ("area_full_storage_km2", "Reservoir area at full storage (Km2)", "km2"),
    ("dam_height_m", "Dam Height  (m)", "m"),
    ("min_operating_level_masl", "Reservoir min operating level (masl)", "masl"),
    ("max_operating_level_masl", "Reservoir max operating level (masl)", "masl"),
    ("elevation_masl", "Reservoir elevation  (masl)", "masl"),
    ("latitude", "Coordinates (y)", "deg"),
    ("longitude", "Coordinates (x)", "deg"),
]

HYDRO_ROR_COLUMNS = [
    ("plant_name", "Hydropower plant name", None),
    ("turbine_max_capacity_mw", "Turbine max capacity (MW)", "MW"),
    ("turbine_max_discharge_m3s", "Turbine max capacity (m3/s)", "m3/s"),
    ("min_operating_level_masl", "Min Operating water level for RoR (masl)", "masl"),
    ("max_operating_level_masl", "Max operating water level (masl)", "masl"),
    ("turbine_elevation_masl", "Turbine elevation (masl)", "masl"),
    ("head_m", "Head (m)", "m"),
    ("exists", "Exists (yes / no)", None),
    ("purpose", "Purpose", None),
    ("latitude", "Coordinates (y)", "deg"),
    ("longitude", "Coordinates (x)", "deg"),
]

HYDRO_STORAGE_COLUMNS = [
    ("plant_name", "Hydropower plant name", None),
    ("reservoir_name", "Reservoir name", None),
    ("regulation", "Regulation (Inter-annual / Seasonal)", None),
    ("reservoir_max_volume_mm3", "Reservoir max volume (Mm3)", "Mm3"),
    ("reservoir_active_volume_mm3", "Reservoir active volume (Mm3)", "Mm3"),
    ("dam_height_m", "Dam Height (m)", "m"),
    ("min_operating_level_masl", "Reservoir min operating level (masl)", "masl"),
    ("max_operating_level_masl", "Reservoir max operating level (masl)", "masl"),
    ("turbine_elevation_masl", "Turbine elevation (masl)", "masl"),
    ("turbine_max_capacity_mw", "Turbine max capacity (MW)", "MW"),
    ("turbine_max_discharge_m3s", "Turbine max capacity (m3/s)", "m3/s"),
    ("latitude", "Coordinates (y)", "deg"),
    ("longitude", "Coordinates (x)", "deg"),
]

SUMMARY_COLUMNS = [
    ("number", "No.", None),
    ("plant_name", "Power plant name", None),
    ("regulation_type", "Regulation type", None),
    ("generating_units", "Installed capacity (MW) / Per unit", "units x MW"),
    ("installed_capacity_mw", "Installed capacity (MW) / Total", "MW"),
    ("total_storage", "Reservoir volume (million m3 / million kWh) / Total storage", "million m3 / million kWh"),
    ("live_storage", "Reservoir volume (million m3 / million kWh) / Live storage", "million m3 / million kWh"),
    ("head_used_m", "Head (m) / Used in these calculations", "m"),
    ("head_max_min_m", "Head (m) / MAX/MIN", "m"),
    ("water_level_max_min_m", "Water level (m) / MAX/MIN", "m"),
    ("water_level_catastrophic_m", "Water level (m) / Catastrophic", "m"),
    ("specific_water_m3_per_mw", "Specific cost of water for 1 MW (m3)", "m3"),
    ("rated_discharge_m3s", "Rated discharge (m3)", "m3/s"),
    ("turbine_type", "Turbine / Type", None),
    ("turbine_note", "Turbine / (unlabelled column)", None),
    ("stator_voltage_kv", "Stator voltage (kV)", "kV"),
    ("dam_height_m", "Dam height (m)", "m"),
    ("dam_type", "Dam type", None),
]

SUMMARY_SHEETS = ("HPP Summary", "Stations by Location", "Hydropower plants", "Reservoirs")


def _clean(value: Any) -> Any:
    if isinstance(value, str):
        stripped = value.strip()
        return stripped or None
    if isinstance(value, float) and value.is_integer() and abs(value) < 1e15:
        # Keep integral floats readable; never alter non-integral values.
        return value
    return value


def _trimmed(row: tuple[Any, ...]) -> list[Any]:
    values = [_clean(v) for v in row]
    while values and values[-1] is None:
        values.pop()
    return values


def _columns(spec: list[tuple[str, str, str | None]]) -> list[dict[str, Any]]:
    return [{"key": key, "source_header": header, "unit": unit} for key, header, unit in spec]


def _record(table: str, row_number: int, spec, cells: list[Any], *, section: str,
            record_type: str, sheet: str) -> dict[str, Any]:
    values = {key: (cells[i] if i < len(cells) else None) for i, (key, _, _) in enumerate(spec)}
    extra = {f"column_{i + 1}": cells[i] for i in range(len(spec), len(cells)) if cells[i] is not None}
    record = {"record_id": f"{table}:r{row_number}", "source_sheet": sheet, "source_row": row_number,
              "section": section, "record_type": record_type, "values": values}
    if extra:
        record["extra_cells"] = extra
    return record


def _table(table: str, sheet: str, spec, records: list[dict[str, Any]], sha: str, notes: list[str]) -> dict[str, Any]:
    return {"schema": SCHEMA, "table": table, "source_workbook": "Power Plant Parameters English.xlsx",
            "source_sha256": sha, "source_sheet": sheet, "columns": _columns(spec),
            "layout_notes": notes, "records": records}


def build_stations(ws, sha: str) -> dict[str, Any]:
    rows = list(ws.iter_rows(values_only=True))
    records, section = [], "Hydro Power Plants"
    for index, row in enumerate(rows[1:], start=2):
        cells = _trimmed(row)
        if not cells:
            continue
        name = cells[0]
        others = [c for c in cells[1:] if c is not None]
        if isinstance(name, str) and name.lower().startswith("total"):
            record_type = "aggregate_total"
        elif name is not None and not others:
            record_type, section = "section_header", name
        else:
            record_type = "plant"
        records.append(_record("stations", index, STATION_COLUMNS, cells, section=section,
                               record_type=record_type, sheet=ws.title))
    return _table("stations", ws.title, STATION_COLUMNS, records, sha, [
        "Row 1 is the column header.",
        "Rows whose only value is a name are section headings; rows named 'Total ...' are aggregates.",
        "The first section has no heading row in the workbook and is labelled 'Hydro Power Plants'.",
    ])


def build_reservoirs(ws, sha: str) -> dict[str, Any]:
    rows = list(ws.iter_rows(values_only=True))
    records = []
    for index, row in enumerate(rows[1:], start=2):
        cells = _trimmed(row)
        if not cells:
            continue
        records.append(_record("reservoirs", index, RESERVOIR_COLUMNS, cells, section="Reservoirs",
                               record_type="reservoir", sheet=ws.title))
    return _table("reservoirs", ws.title, RESERVOIR_COLUMNS, records, sha, ["Row 1 is the column header."])


def build_hydropower(ws, sha: str) -> dict[str, Any]:
    rows = list(ws.iter_rows(values_only=True))
    records: list[dict[str, Any]] = []
    section, spec, in_data = None, None, False
    for index, row in enumerate(rows, start=1):
        cells = _trimmed(row)
        if not cells:
            continue
        first = cells[0]
        if isinstance(first, str) and first.lower().startswith("hydropower plants (") and len(cells) == 1:
            section = first
            spec = HYDRO_STORAGE_COLUMNS if "storage" in first.lower() else HYDRO_ROR_COLUMNS
            in_data = False
            continue
        if first == "Hydropower plant name":
            in_data = True
            continue
        if not in_data or spec is None:
            continue  # group-header row above the unit row
        table = "hydropower_storage" if spec is HYDRO_STORAGE_COLUMNS else "hydropower_ror"
        records.append(_record(table, index, spec, cells, section=section or "", record_type="plant",
                               sheet=ws.title))
    for record in records:
        record["record_id"] = "hydropower:" + record["record_id"].split(":", 1)[1]
    return {"schema": SCHEMA, "table": "hydropower", "source_workbook": "Power Plant Parameters English.xlsx",
            "source_sha256": sha, "source_sheet": ws.title,
            "sections": {
                "Hydropower plants (RoR)": {"columns": _columns(HYDRO_ROR_COLUMNS)},
                "Hydropower plants (Storage)": {"columns": _columns(HYDRO_STORAGE_COLUMNS)},
            },
            "layout_notes": [
                "The sheet holds two sections with different column layouts, each with a two-row header "
                "(group header + unit row).",
                "Header rows are not data records. A one-header export misreads the storage section's "
                "group-header row as a plant row without a name.",
            ],
            "records": records}


def build_summary(ws, sha: str) -> dict[str, Any]:
    rows = list(ws.iter_rows(values_only=True))
    records = []
    for index, row in enumerate(rows[2:], start=3):
        cells = _trimmed(row)
        if not cells:
            continue
        records.append(_record("hpp_summary", index, SUMMARY_COLUMNS, cells, section="HPP Summary",
                               record_type="plant", sheet=ws.title))
    return _table("hpp_summary", ws.title, SUMMARY_COLUMNS, records, sha,
                  ["Rows 1-2 are a two-row header (group + sub-column)."])


def build_detail_sheets(workbook, sha: str) -> dict[str, Any]:
    sheets = {}
    for name in workbook.sheetnames:
        if name in SUMMARY_SHEETS:
            continue
        rows = []
        for index, row in enumerate(workbook[name].iter_rows(values_only=True), start=1):
            cells = _trimmed(row)
            if cells:
                rows.append({"source_row": index, "cells": cells})
        sheets[name] = {"record_id": f"sheet:{name}", "source_sheet": name, "rows": rows}
    return {"schema": SCHEMA, "table": "plant_detail_sheets", "source_workbook": "Power Plant Parameters English.xlsx",
            "source_sha256": sha, "layout_notes": ["One key/value technical sheet per plant; empty rows omitted."],
            "sheets": sheets}


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--out", type=Path, default=Path("data/reference/plant_parameters"))
    args = parser.parse_args(argv)
    import openpyxl

    raw = args.workbook.read_bytes()
    sha = hashlib.sha256(raw).hexdigest()
    workbook = openpyxl.load_workbook(args.workbook, read_only=True, data_only=True)
    outputs = {
        "stations_by_location.json": build_stations(workbook["Stations by Location"], sha),
        "reservoirs.json": build_reservoirs(workbook["Reservoirs"], sha),
        "hydropower_assets.json": build_hydropower(workbook["Hydropower plants"], sha),
        "hpp_summary.json": build_summary(workbook["HPP Summary"], sha),
        "plant_detail_sheets.json": build_detail_sheets(workbook, sha),
    }
    args.out.mkdir(parents=True, exist_ok=True)
    for name, payload in outputs.items():
        (args.out / name).write_text(json.dumps(payload, ensure_ascii=False, indent=1, default=str) + "\n",
                                     encoding="utf-8")
    manifest = {
        "schema": "powersim.reference.manifest.v2",
        "source_file": "Power Plant Parameters English.xlsx",
        "sha256": sha,
        "generator": "scripts/build_reference_catalogue.py",
        "sheet_count": len(workbook.sheetnames),
        "sheets": workbook.sheetnames,
        "record_counts": {
            "stations_by_location": len(outputs["stations_by_location.json"]["records"]),
            "stations_by_location_plants": sum(r["record_type"] == "plant"
                                               for r in outputs["stations_by_location.json"]["records"]),
            "reservoirs": len(outputs["reservoirs.json"]["records"]),
            "hydropower_assets": len(outputs["hydropower_assets.json"]["records"]),
            "hpp_summary": len(outputs["hpp_summary.json"]["records"]),
            "detail_sheets": len(outputs["plant_detail_sheets.json"]["sheets"]),
        },
        "important_note": ("This is a reference-data import. Capacity, unit and operating values must be "
                           "validated against the raw workbook and approved study inputs before they are "
                           "used in a calculation."),
    }
    (args.out / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n",
                                            encoding="utf-8")
    print(json.dumps(manifest["record_counts"]))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
