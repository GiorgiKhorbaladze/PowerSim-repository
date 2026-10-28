"""Bulk asset-parameter Excel workbook: template, checks, apply and audit."""
from __future__ import annotations

import io
import json
from pathlib import Path

import openpyxl
import pytest

from powersim.contracts import ProjectContract
from powersim.parameter_workbook import WorkbookRejected, apply_workbook, build_template, read_workbook
from powersim.reference import ReferenceCatalogue
from powersim.reference.importer import apply_import
from powersim.validation.readiness import diagnose_project

ROOT = Path(__file__).resolve().parents[1]
GARDABANI = {"record_id": "stations:r140", "overrides": {
    "capacity_max_mw": {"value": 231.2, "reason": "nameplate"}, "fuel_type": {"value": "natural gas", "reason": "CCGT"}}}


def _compact() -> ProjectContract:
    return ProjectContract.model_validate_json((ROOT / "samples/application_compact_demo.json").read_text(encoding="utf-8"))


def _edit(project: ProjectContract, rows: dict[str, dict], profiles: dict[str, list] | None = None) -> bytes:
    workbook = openpyxl.load_workbook(io.BytesIO(build_template(project)))
    sheet = workbook["Assets"]
    header = [c.value for c in sheet[1]]
    for row in sheet.iter_rows(min_row=3):
        values = rows.get(row[0].value)
        for column, value in (values or {}).items():
            row[header.index(column)].value = value
    for asset_id, values in rows.items():
        if asset_id not in {r[0].value for r in sheet.iter_rows(min_row=3)}:
            sheet.append([values.get(column, asset_id if column == "asset_id" else None) for column in header])
    if profiles:
        sheet = workbook["Profiles"]
        for profile_id, values in profiles.items():
            column = sheet.max_column + 1
            sheet.cell(1, column, profile_id)
            for index, value in enumerate(values, start=2):
                sheet.cell(index, column, value)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_template_lists_assets_and_profiles_and_round_trips_without_changes():
    project = _compact()
    workbook = openpyxl.load_workbook(io.BytesIO(build_template(project)))
    assert workbook.sheetnames == ["Instructions", "Assets", "Profiles"]
    ids = [row[0].value for row in workbook["Assets"].iter_rows(min_row=3)]
    assert ids == ["thermal", "import"]
    assert [c.value for c in workbook["Profiles"][1]] == ["period", "demand"]
    report = read_workbook(project, build_template(project))
    assert report["errors"] == 0 and report["changes"] == [] and not report["applicable"]


def test_parameters_are_applied_as_an_audited_new_revision():
    project = _compact()
    content = _edit(project, {"thermal": {"vom_per_mwh": 12.5, "heat_rate_gj_per_mwh": "9,5", "fuel_price_usd_per_mmbtu": 4,
                                          "fuel_type": "natural gas", "capacity_max_mw": 35},
                              "import": {"vom_per_mwh": 80}})
    report = read_workbook(project, content)
    assert report["errors"] == 0 and report["applicable"]
    changes = {(c["asset_id"], c["field"]): (c["old"], c["new"]) for c in report["changes"]}
    assert changes[("thermal", "vom_per_mwh")] == (10, 12.5)
    assert changes[("thermal", "heat_rate_gj_per_mwh")] == (0, 9.5)  # comma decimal accepted
    assert changes[("import", "vom_per_mwh")] == (100, 80)
    updated, _ = apply_workbook(project, content, file_name="params.xlsx", imported_by="tester")
    thermal = next(a for a in updated.assets if a.id == "thermal")
    assert thermal.capacity_max_mw == 35 and thermal.legacy_extensions["pmax"] == 35
    assert thermal.legacy_extensions["vom"] == 12.5 and thermal.legacy_extensions["fuel_type"] == "natural gas"
    assert updated.version.revision == 2 and updated.version.parent_fingerprint == project.fingerprint()
    record = updated.metadata["parameter_imports"][-1]
    assert record["file_name"] == "params.xlsx" and record["imported_by"] == "tester"
    assert record["file_sha256"].startswith("sha256:") and len(record["changes"]) == len(report["changes"])
    assert diagnose_project(updated.model_dump(mode="json"))["run_ready"]


@pytest.mark.parametrize("rows,code", [
    ({"ghost": {"vom_per_mwh": 1}}, "unknown_asset"),
    ({"thermal": {"vom_per_mwh": -1}}, "invalid_number"),
    ({"thermal": {"vom_per_mwh": "cheap"}}, "invalid_number"),
    ({"thermal": {"kind": "nuclear"}}, "invalid_kind"),
    ({"thermal": {"enabled": "maybe"}}, "invalid_boolean"),
    ({"thermal": {"capacity_min_mw": 40}}, "capacity_order"),
    ({"thermal": {"availability_profile": "nope"}}, "unknown_profile"),
    ({"thermal": {"availability_profile": "demand"}}, "profile_not_applicable"),
])
def test_invalid_cells_are_errors_with_cell_reference(rows, code):
    project = _compact()
    report = read_workbook(project, _edit(project, rows))
    errors = [i for i in report["issues"] if i["severity"] == "error"]
    assert code in {i["code"] for i in errors}
    assert all(i["path"].startswith("Assets!") for i in errors)
    with pytest.raises(WorkbookRejected):
        apply_workbook(project, _edit(project, rows))


def test_profile_length_and_file_type_are_checked():
    project = _compact()
    report = read_workbook(project, _edit(project, {}, {"cf": [0.5, 0.5, 0.5]}))
    assert "profile_length_mismatch" in {i["code"] for i in report["issues"]}
    assert read_workbook(project, b"not a workbook")["issues"][0]["code"] == "not_an_xlsx_workbook"
    with pytest.raises(WorkbookRejected):
        apply_workbook(project, build_template(project))  # no changes


def test_reference_asset_completed_and_reviewed_through_the_workbook():
    catalogue = ReferenceCatalogue()
    project, _ = apply_import(_compact(), catalogue, [GARDABANI, {"record_id": "stations:r2"}], acknowledge_warnings=True)
    # enabling without review is a warning here and a run blocker in validation
    blocked = read_workbook(project, _edit(project, {"ref-gardabani-tpp-1-ccpp-1": {"enabled": True}}))
    assert "reference_asset_unreviewed" in {i["code"] for i in blocked["issues"]}
    no_note = read_workbook(project, _edit(project, {"ref-gardabani-tpp-1-ccpp-1": {"reviewed": "yes"}}))
    assert "review_note_required" in {i["code"] for i in no_note["issues"] if i["severity"] == "error"}

    content = _edit(project, {
        "ref-gardabani-tpp-1-ccpp-1": {"enabled": True, "heat_rate_gj_per_mwh": 7.2, "fuel_price_usd_per_mmbtu": 6,
                                       "vom_per_mwh": 3, "reviewed": "yes", "review_note": "checked vs GSE register"},
        "ref-abhesi-hpp": {"enabled": True, "availability_profile": "abhesi_cf", "reviewed": "yes",
                           "review_note": "flow record 2020-2024"},
    }, {"abhesi_cf": [0.4, 0.5, 0.6, 0.5]})
    report = read_workbook(project, content)
    assert report["errors"] == 0 and report["profiles"] == [{"id": "abhesi_cf", "values": 4, "action": "added"}]
    updated, _ = apply_workbook(project, content, imported_by="reviewer")
    gas = next(a for a in updated.assets if a.id == "ref-gardabani-tpp-1-ccpp-1")
    details = gas.provenance.details
    assert gas.enabled and details["review_status"] == "user_reviewed" and details["reviewed_by"] == "reviewer"
    assert details["review_note"] == "checked vs GSE register"
    assert details["field_origins"]["heat_rate_gj_per_mwh"]["origin"] == "parameter_workbook"
    ror = next(a for a in updated.assets if a.id == "ref-abhesi-hpp")
    assert ror.profile_references == ("abhesi_cf",)
    profile = next(p for p in updated.profiles if p.id == "abhesi_cf")
    assert profile.unit == "per-unit" and profile.provenance.source == "parameter_workbook"
    diagnostics = diagnose_project(updated.model_dump(mode="json"))
    assert diagnostics["run_ready"], [i for i in diagnostics["issues"] if i["severity"] == "error"]


def test_parameter_workbook_http_routes(tmp_path):
    pytest.importorskip("httpx")
    from fastapi.testclient import TestClient
    from powersim.application import ApplicationService, build_fastapi_app
    from powersim.platform import RunManager
    service = ApplicationService(RunManager(tmp_path), asynchronous=False)
    client = TestClient(build_fastapi_app(service))
    assert client.get("/projects/application-compact-demo/parameter-workbook").status_code == 404
    client.post("/projects/save", json=json.loads((ROOT / "samples/application_compact_demo.json").read_text()))
    template = client.get("/projects/application-compact-demo/parameter-workbook")
    assert template.status_code == 200 and template.headers["content-type"].startswith("application/vnd.openxml")
    project = _compact()
    content = _edit(project, {"import": {"vom_per_mwh": 70}})
    check = client.post("/projects/application-compact-demo/parameter-workbook", content=content).json()
    assert check["applied"] is False and check["applicable"] and check["changes"][0]["new"] == 70
    assert client.get("/projects/application-compact-demo").json()["project"]["assets"][1]["legacy_extensions"]["vom"] == 100
    applied = client.post("/projects/application-compact-demo/parameter-workbook?apply=true&file_name=p.xlsx",
                          content=content).json()
    assert applied["applied"] and applied["revision"] == 2 and applied["run_ready"]
    assert client.get("/projects/application-compact-demo").json()["project"]["assets"][1]["legacy_extensions"]["vom"] == 70
    bad = client.post("/projects/application-compact-demo/parameter-workbook?apply=true",
                      content=_edit(project, {"import": {"vom_per_mwh": -5}}))
    assert bad.status_code == 422 and bad.json()["detail"]["code"] == "workbook_rejected"
    assert client.post("/projects/application-compact-demo/parameter-workbook", content=b"").status_code == 422
